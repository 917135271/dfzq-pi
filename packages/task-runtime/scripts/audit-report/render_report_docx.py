from __future__ import annotations

import argparse
import json
import os
import re
import tempfile
import unicodedata
import zipfile
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.shared import Twips
from docx.text.paragraph import Paragraph
from lxml import etree

WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def first_text_run(paragraph: Any) -> Any:
    run = next((item for item in paragraph.runs if item.text.strip()), None)
    if run is None:
        raise ValueError(f"Template prototype has no formatted run: {paragraph.text!r}")
    return run


def clone_paragraph(document: DocumentObject, prototype: Any, text: str) -> Any:
    if prototype is None:
        raise ValueError(f"Template does not provide a prototype for: {text[:40]!r}")
    paragraph = document.add_paragraph()
    try:
        paragraph.style = prototype.style
    except (KeyError, ValueError):
        pass
    if prototype._p.pPr is not None:
        if paragraph._p.pPr is not None:
            paragraph._p.remove(paragraph._p.pPr)
        paragraph._p.insert(0, deepcopy(prototype._p.pPr))
    run = paragraph.add_run(text)
    source_run = first_text_run(prototype)
    if source_run._r.rPr is not None:
        run._r.insert(0, deepcopy(source_run._r.rPr))
    return paragraph


def find_prototypes(document: DocumentObject) -> dict[str, Any]:
    paragraphs = [item for item in document.paragraphs if item.text.strip()]
    if len(paragraphs) < 4:
        raise ValueError("Template must contain formatted title, body, closing and date prototypes")

    def find(pattern: str) -> Any | None:
        return next((item for item in paragraphs if re.search(pattern, item.text.strip())), None)

    date = find(r"^\d{4}年.+日$")
    if date is None:
        raise ValueError("Template does not contain a formatted report-date prototype")
    date_index = paragraphs.index(date)
    closing = paragraphs[date_index - 1] if date_index > 0 else None
    body = find(r"^(?:按照审计工作安排|根据.+委托)")
    prototypes = {
        "title-1": paragraphs[0],
        "title-2": paragraphs[1] if len(paragraphs) > 1 else paragraphs[0],
        "addressee": find(r"[：:]$"),
        "body": body,
        "section": find(r"^[一二三四五六七八九十]+、"),
        "subsection": find(r"^（[一二三四五六七八九十]+）"),
        "finding-title": find(r"^\d+[.．、]"),
        "note": find(r"^备注："),
        "table-title": find(r"^表\d+"),
        "closing": closing,
        "date": date,
    }
    for key in ("body", "section", "closing", "date"):
        if prototypes[key] is None:
            raise ValueError(f"Template does not contain required prototype: {key}")
    return prototypes


def clear_body(document: DocumentObject) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clone_cell_format(target: Any, source: Any, value: Any) -> None:
    target.text = ""
    if source._tc.tcPr is not None:
        if target._tc.tcPr is not None:
            target._tc.remove(target._tc.tcPr)
        target._tc.insert(0, deepcopy(source._tc.tcPr))
    source_paragraph = source.paragraphs[0]
    target_paragraph = target.paragraphs[0]
    if source_paragraph._p.pPr is not None:
        target_paragraph._p.insert(0, deepcopy(source_paragraph._p.pPr))
    run = target_paragraph.add_run(str(value))
    # Some official table prototypes intentionally keep data cells empty. In
    # that case paragraph/cell properties still carry the layout, while there
    # is no non-empty run whose character properties can be copied.
    source_run = next((item for item in source_paragraph.runs if item.text.strip()), None)
    if source_run is None:
        source_run = next(iter(source_paragraph.runs), None)
    if source_run is not None and source_run._r.rPr is not None:
        run._r.insert(0, deepcopy(source_run._r.rPr))
    numeric_text = str(value).strip().replace("−", "-")
    if re.fullmatch(r"[-+]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?%?", numeric_text):
        # Values arrive as formatted strings as well as numbers. Template
        # samples may be red losses; never inherit that data-dependent color.
        amount = Decimal(numeric_text.replace(",", "").rstrip("%"))
        run.font.color.rgb = RGBColor(255, 0, 0) if amount < 0 else RGBColor(0, 0, 0)


def add_table(document: DocumentObject, data: dict[str, Any], title_prototype: Any, prototype: Any) -> None:
    if isinstance(prototype, list):
        # Capture each table with its own caption before clearing the body.
        # Financial, performance and ranking tables have distinct row styles.
        wanted = re.sub(r"^表\s*\d+\s*[：:.．]?\s*", "", data["title"]).strip()
        selected = next((pair for pair in prototype if pair[0] is not None and
                         re.sub(r"^表\s*\d+\s*[：:.．]?\s*", "", pair[0].text.split("单位")[0]).strip() == wanted), None)
        if selected is None:
            selected = prototype[0] if prototype else (title_prototype, None)
        title_prototype, prototype = selected[0] or title_prototype, selected[1]
    if prototype is None:
        raise ValueError(f"Template has no table prototype for {data['tableId']}")
    if data.get("unit"):
        separator = re.search(r"(\s+)单位[：:]", title_prototype.text)
        spacing = separator.group(1) if separator else "\t"
    caption = clone_paragraph(document, title_prototype, data["title"])
    caption.paragraph_format.keep_with_next = True
    if data.get("unit"):
        # Preserve the formatting at the spacing/unit positions, not the
        # first (often underlined) title run for the entire caption.
        unit_position = title_prototype.text.find("单位")
        spacing_position = separator.start(1) if separator else unit_position
        for text, position in ((spacing, spacing_position), (f"单位：{data['unit']}", unit_position)):
            offset = 0
            source_run = first_text_run(title_prototype)
            for candidate in title_prototype.runs:
                if offset <= position < offset + len(candidate.text):
                    source_run = candidate
                    break
                offset += len(candidate.text)
            run = caption.add_run(text)
            if source_run._r.rPr is not None:
                run._r.insert(0, deepcopy(source_run._r.rPr))
    headers = data["headers"]
    table = document.add_table(rows=1, cols=len(headers))
    if prototype._tbl.tblPr is not None:
        table._tbl.remove(table._tbl.tblPr)
        table._tbl.insert(0, deepcopy(prototype._tbl.tblPr))
    # Reuse the official column grid only when the generated table has the
    # same number of columns. Copying a five-column prototype grid into a
    # three-column period table leaves visible blank columns in Word.
    if prototype._tbl.tblGrid is not None and len(prototype.columns) == len(headers):
        table._tbl.remove(table._tbl.tblGrid)
        table._tbl.insert(1, deepcopy(prototype._tbl.tblGrid))
    elif prototype._tbl.tblGrid is not None and len(headers) > 1:
        # Preserve the label column and total template width when periods change.
        # The cell widths must agree with the new grid, not old percentage widths.
        source_widths = [int(column.get(qn("w:w"))) for column in prototype._tbl.tblGrid]
        remaining = sum(source_widths) - source_widths[0]
        width, extra = divmod(remaining, len(headers) - 1)
        widths = [source_widths[0]] + [width + (1 if index < extra else 0) for index in range(len(headers) - 1)]
        for column, value in zip(table._tbl.tblGrid, widths):
            column.set(qn("w:w"), str(value))
    for index, header in enumerate(headers):
        clone_cell_format(table.rows[0].cells[index], prototype.rows[0].cells[min(index, len(prototype.columns) - 1)], header)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    def label(value: Any) -> str:
        return re.sub(r"\s+", "", str(value)).replace(":", "：")

    fallback_row = prototype.rows[min(1, len(prototype.rows) - 1)]
    for values in data["rows"]:
        source_row = next((item for item in prototype.rows[1:] if values and
                           label(item.cells[0].text) == label(values[0])), fallback_row)
        row = table.add_row()
        if source_row._tr.trPr is not None:
            row._tr.insert(0, deepcopy(source_row._tr.trPr))
        for index, value in enumerate(values):
            source_cell = source_row.cells[min(index, len(prototype.columns) - 1)]
            if index == 0 and label(source_cell.text) == label(value):
                value = source_cell.text  # retain label spacing/indentation only, never sample numbers
            clone_cell_format(row.cells[index], source_cell, value)
    # Reserve space for numeric tokens at the retained font size. Some renderers
    # still wrap noWrap cells at their preferred width, so noWrap alone is not
    # enough. Borrow from the label column, which can safely grow vertically.
    widths = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid]
    for column_index in range(1, len(widths)):
        for row in table.rows[1:]:
            cell = row.cells[column_index]
            if re.fullmatch(r"[-+−]?\d[\d,]*(?:\.\d+)?%?", cell.text.strip()):
                run = first_text_run(cell.paragraphs[0])
                size = run.font.size or cell.paragraphs[0].style.font.size
                if size is not None:
                    # Conservative ASCII glyph allowance plus template cell padding.
                    padding = 216
                    margins = table._tbl.tblPr.find(qn("w:tblCellMar"))
                    if margins is not None:
                        padding = sum(int(margins.find(qn(f"w:{side}")).get(qn("w:w"), "108"))
                                      if margins.find(qn(f"w:{side}")) is not None else 108
                                      for side in ("left", "right"))
                    widths[column_index] = max(widths[column_index], round(len(cell.text.strip()) * size.pt * 12 + padding))
    total = sum(int(column.get(qn("w:w"))) for column in table._tbl.tblGrid)
    widths[0] = total - sum(widths[1:])
    if widths[0] < total / 5:
        raise ValueError(f"Table {data['tableId']} has too many periods or numeric digits for the template width")
    for column, width in zip(table._tbl.tblGrid, widths):
        column.set(qn("w:w"), str(width))
    # The widths above are the actual layout contract. Leaving the retained
    # percentage preferred width and AutoFit enabled lets Word redistribute
    # columns again (e.g. for a longer period header), invalidating both the
    # numeric-width reservation and the label-indent calculation below.
    table.autofit = False
    preferred_width = table._tbl.tblPr.find(qn("w:tblW"))
    if preferred_width is not None:
        preferred_width.set(qn("w:type"), "dxa")
        preferred_width.set(qn("w:w"), str(sum(widths)))
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            # Repeating a header does not keep it with the first data row.
            # Chain only the header paragraphs, not the entire data table.
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = bool(data["rows"])
        if row._tr.trPr is not None:
            for height in row._tr.trPr.findall(qn("w:trHeight")):
                if height.get(qn("w:hRule")) == "exact":
                    height.set(qn("w:hRule"), "atLeast")
        for column_index, (cell, column) in enumerate(zip(row.cells, table._tbl.tblGrid)):
            cell.width = Twips(int(column.get(qn("w:w"))))
            if row_index > 0 and column_index == 0:
                # A fixed first-line indent from a wider template column can
                # strand the final character of an otherwise fitting label.
                paragraph = cell.paragraphs[0]
                indent = paragraph._p.pPr.find(qn("w:ind")) if paragraph._p.pPr is not None else None
                run = next((item for item in paragraph.runs if item.text.strip()), None)
                size = (run.font.size or paragraph.style.font.size) if run is not None else None
                if indent is not None and size is not None:
                    chars = indent.get(qn("w:firstLineChars"))
                    current = round(int(chars) * size.pt / 5) if chars is not None else int(indent.get(qn("w:firstLine"), "0"))
                    units = sum(1 if unicodedata.east_asian_width(char) in ("W", "F") else 0.6 for char in cell.text)
                    maximum = max(0, int(column.get(qn("w:w"))) - 216 - round(units * size.pt * 20))
                    if current > maximum:
                        indent.attrib.pop(qn("w:firstLineChars"), None)
                        indent.set(qn("w:firstLine"), str(maximum))
            for no_wrap in cell._tc.tcPr.findall(qn("w:noWrap")):
                cell._tc.tcPr.remove(no_wrap)
            if row_index > 0 and column_index > 0 and re.fullmatch(r"[-+−]?\d[\d,]*(?:\.\d+)?%?", cell.text.strip()):
                cell._tc.tcPr.append(OxmlElement("w:noWrap"))


def add_paragraph(document: DocumentObject, item: dict[str, Any], prototypes: dict[str, Any]) -> None:
    if item["paragraphId"].endswith("-title"):
        key = "finding-title"
    elif item["text"].startswith("备注："):
        key = "note"
    else:
        key = "body"
    # A report template may not contain an example finding title or note when
    # that particular sample has no findings. Preserve report generation by
    # falling back to the template's body style for these optional prototypes.
    paragraph = clone_paragraph(document, prototypes.get(key) or prototypes["body"], item["text"])
    if key == "finding-title":
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
    else:
        # Template samples can disable widow control for their original length.
        # Generated facts have variable lengths: retain normal paragraph splits,
        # but do not strand a single opening or closing line across pages.
        paragraph.paragraph_format.widow_control = True


def add_draft(document: DocumentObject, draft: dict[str, Any], prototypes: dict[str, Any], table_prototype: Any) -> None:
    for index, title in enumerate(draft["titleLines"]):
        prototype = prototypes[f"title-{min(index + 1, 2)}"]
        display_title = title
        if index == 0 and draft.get("reportType") == "turnover" and "营业部" in title:
            # Keep a long position name (including its parenthetical qualifier)
            # together. This is a layout break within the same saved title node,
            # not a rewritten title or a smaller font.
            size = first_text_run(prototype).font.size or prototype.style.font.size
            section = document.sections[-1]
            available = (section.page_width - section.left_margin - section.right_margin) / 12700
            parts = title.partition("营业部")
            organization, position = parts[0] + parts[1], parts[2]
            widths = [sum(1 if unicodedata.east_asian_width(char) in ("W", "F") else 0.6
                          for char in text) * size.pt for text in (title, organization, position)] if size is not None else []
            if position and widths and widths[0] > available and max(widths[1:]) <= available:
                display_title = organization + "\n" + position
        paragraph = clone_paragraph(document, prototype, display_title)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
    if draft.get("addressee"):
        clone_paragraph(document, prototypes["addressee"], draft["addressee"])
    clone_paragraph(document, prototypes["body"], draft["introduction"]["text"])
    for section in draft["sections"]:
        heading = clone_paragraph(document, prototypes["section"], section["heading"])
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.keep_together = True
        for item in section["paragraphs"]:
            add_paragraph(document, item, prototypes)
        for table in section["tables"]:
            add_table(document, table, prototypes["table-title"], table_prototype)
            for note in table.get("notes", []):
                clone_paragraph(document, prototypes.get("note") or prototypes["body"], note)
        for subsection in section["subsections"]:
            heading = clone_paragraph(document, prototypes["subsection"], subsection["heading"])
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.keep_together = True
            split = subsection.get("tablesAfterParagraphCount", len(subsection["paragraphs"]))
            for index, item in enumerate(subsection["paragraphs"]):
                if index == split:
                    for table in subsection.get("tables", []):
                        add_table(document, table, prototypes["table-title"], table_prototype)
                        for note in table.get("notes", []):
                            clone_paragraph(document, prototypes.get("note") or prototypes["body"], note)
                add_paragraph(document, item, prototypes)
            if split >= len(subsection["paragraphs"]):
                for table in subsection.get("tables", []):
                    add_table(document, table, prototypes["table-title"], table_prototype)
                    for note in table.get("notes", []):
                        clone_paragraph(document, prototypes.get("note") or prototypes["body"], note)
        for item in section.get("closingParagraphs", []):
            add_paragraph(document, item, prototypes)
    closing_prefix = re.match(r"^\s*", prototypes["closing"].text).group(0)
    clone_paragraph(document, prototypes["closing"], closing_prefix + draft["closingOrganization"])
    # Some retained templates use leading spaces in a centered paragraph to
    # position the date under the right-aligned organization. They are layout,
    # not part of the business date, and must survive text substitution.
    date_prefix = re.match(r"^\s*", prototypes["date"].text).group(0)
    clone_paragraph(document, prototypes["date"], date_prefix + draft["reportDate"])
    if draft["status"] == "needs-input":
        clone_paragraph(document, prototypes["body"], "【数据未就绪：本草稿不得作为正式审计报告】")


def sanitize_review_markup(path: Path) -> None:
    removable_parts = {"word/comments.xml", "word/commentsExtended.xml", "word/commentsExtensible.xml", "word/commentsIds.xml", "word/people.xml"}
    removable_tags = {f"{{{WORD_NAMESPACE}}}{name}" for name in ("commentRangeStart", "commentRangeEnd", "commentReference", "del", "moveFrom", "pPrChange", "rPrChange", "tblPrChange", "tblGridChange", "trPrChange", "tcPrChange", "sectPrChange", "numPrChange")}
    accepted_tags = {f"{{{WORD_NAMESPACE}}}{name}" for name in ("ins", "moveTo")}
    with tempfile.NamedTemporaryFile(suffix=".docx", dir=path.parent, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.filename in removable_parts:
                    continue
                payload = source.read(info.filename)
                if info.filename.endswith((".xml", ".rels")):
                    root = etree.fromstring(payload)
                    changed = False
                    for element in list(root.iter()):
                        parent = element.getparent()
                        if parent is None:
                            continue
                        if element.tag in removable_tags:
                            parent.remove(element)
                            changed = True
                        elif element.tag in accepted_tags:
                            position = parent.index(element)
                            for child in list(element):
                                parent.insert(position, child)
                                position += 1
                            parent.remove(element)
                            changed = True
                        elif any(token in element.get("Target", "") or token in element.get("PartName", "") for token in ("comments", "people")):
                            parent.remove(element)
                            changed = True
                    if changed:
                        payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                target.writestr(info, payload)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Render ReportDraft JSON by inheriting an external DOCX template")
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--draft", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    draft = json.loads(args.draft.read_text(encoding="utf-8"))
    document = Document(str(args.template))
    prototypes = find_prototypes(document)
    table_prototype = []
    for table in document.tables:
        previous = table._tbl.getprevious()
        caption = Paragraph(previous, document) if previous is not None and previous.tag == qn("w:p") else None
        table_prototype.append((caption, table))
    clear_body(document)
    add_draft(document, draft, prototypes, table_prototype)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(args.output))
    sanitize_review_markup(args.output)


if __name__ == "__main__":
    main()
