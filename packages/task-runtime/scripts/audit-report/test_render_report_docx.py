from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from render_report_docx import clone_paragraph, add_table, add_draft, find_prototypes, clear_body


class RenderReportDocxTest(unittest.TestCase):
    def test_distinct_table_rows_and_caption_runs(self) -> None:
        document = Document()
        financial_title = document.add_paragraph("表1：财务指标")
        ranking_title = document.add_paragraph()
        ranking_title.add_run("表3：指标排名情况").underline = True
        ranking_title.add_run("        ").underline = False
        ranking_title.add_run("单位：名").underline = False
        financial = document.add_table(rows=2, cols=2)
        financial.cell(1, 0).text = "收入"
        ranking = document.add_table(rows=3, cols=2)
        for index, name in enumerate(("指标", "收入", "    其中：产品")):
            ranking.cell(index, 0).text = name
            ranking.cell(index, 1).text = "999"
        ranking.cell(2, 0).paragraphs[0].runs[0].italic = True
        ranking.cell(2, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_table(document, {"tableId": "rank", "title": "表3：指标排名情况", "unit": "名", "headers": ["指标", "排名"], "rows": [["其中：产品", 2], ["收入", 3]]}, financial_title, [(financial_title, financial), (ranking_title, ranking)])
        result = document.tables[-1]
        self.assertEqual(result.cell(1, 0).text, "    其中：产品")
        self.assertTrue(result.cell(1, 0).paragraphs[0].runs[-1].italic)
        self.assertEqual(result.cell(1, 1).text, "2")
        self.assertEqual(result.cell(1, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertTrue(document.paragraphs[-1].runs[0].underline)
        self.assertFalse(document.paragraphs[-1].runs[1].underline)
        self.assertFalse(document.paragraphs[-1].runs[2].underline)

    def test_template_spacing_and_single_property_nodes(self) -> None:
        template = Document()
        for text in ("报告标题", "报告副标题", "按照审计工作安排，正文", "一、情况", "表1：指标                    单位：万元", "公司", "                                    2026年X月X日"):
            template.add_paragraph(text, style="Body Text")
        template.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        prototypes = find_prototypes(template)
        table = template.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "指标"
        table.cell(0, 1).text = "数值"
        table.cell(1, 0).text = "收入"
        table.cell(1, 1).text = "1"
        clear_body(template)
        add_draft(template, {"titleLines": ["报告"], "introduction": {"text": "正文"}, "sections": [], "closingOrganization": "公司", "reportDate": "2026年9月11日", "status": "ready-for-review"}, prototypes, table)
        self.assertEqual(template.paragraphs[-1].text, "                                    2026年9月11日")
        self.assertEqual(template.paragraphs[-1].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        for paragraph in template.paragraphs:
            self.assertEqual(len(paragraph._p.xpath("./w:pPr")), 1)
        add_table(template, {"tableId": "t", "title": "表1：指标", "unit": "万元", "headers": ["指标", "数值"], "rows": [["收入", 2]]}, prototypes["table-title"], table)
        self.assertEqual(template.paragraphs[-1].text, "表1：指标                    单位：万元")
        for row in template.tables[-1].rows:
            for cell in row.cells:
                self.assertEqual(len(cell._tc.xpath("./w:tcPr")), 1)

    def test_inherits_external_template_and_writes_draft(self) -> None:
        script = Path(__file__).with_name("render_report_docx.py")
        with tempfile.TemporaryDirectory() as raw_dir:
            root = Path(raw_dir)
            template_path = root / "template.docx"
            draft_path = root / "draft.json"
            output_path = root / "output.docx"

            template = Document()
            for text, size in (
                ("XX证券营业部", 22),
                ("常规审计报告", 22),
                ("XX证券营业部：", 16),
                ("按照审计工作安排，模板正文。", 16),
                ("一、审计项目基本情况", 16),
                ("（一）经营情况", 16),
                ("1.问题标题", 16),
                ("备注：模板说明", 10.5),
                ("表1 经营情况", 11),
                ("模板落款机构", 16),
                ("2026年1月1日", 16),
            ):
                paragraph = template.add_paragraph()
                run = paragraph.add_run(text)
                run.font.size = Pt(size)
            template.save(template_path)

            draft = {
                "taskId": "TASK-001",
                "reportType": "regular",
                "templateId": "regular-v1",
                "templateVersion": "1",
                "titleLines": ["测试证券营业部", "常规审计报告"],
                "addressee": "测试证券营业部：",
                "introduction": {"paragraphId": "intro", "text": "按照计划完成审计。", "evidenceIds": ["E-1"], "requiresHumanReview": False},
                "sections": [{
                    "heading": "一、审计项目基本情况",
                    "paragraphs": [{"paragraphId": "body", "text": "审计期间为2025年度。", "evidenceIds": ["E-1"], "requiresHumanReview": False}],
                    "tables": [],
                    "subsections": [],
                }],
                "closingOrganization": "测试证券公司",
                "reportDate": "2026年1月1日",
                "status": "ready-for-review",
                "blockers": [],
                "warnings": [],
                "allEvidenceIds": ["E-1"],
            }
            draft_path.write_text(json.dumps(draft, ensure_ascii=False), encoding="utf-8")

            subprocess.run(
                [sys.executable, str(script), "--template", str(template_path), "--draft", str(draft_path), "--output", str(output_path)],
                check=True,
            )
            rendered = Document(output_path)
            text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
            self.assertIn("测试证券营业部", text)
            self.assertIn("审计期间为2025年度。", text)
            self.assertIn("测试证券公司", text)
            self.assertEqual(rendered.paragraphs[0].runs[0].font.size, Pt(22))


if __name__ == "__main__":
    unittest.main()
