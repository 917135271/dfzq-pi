"""监督元数据、全量 IR、逐字段引用及现有文档版本读取。"""

import hashlib
import io
import json
from contextlib import contextmanager
from types import SimpleNamespace

import pytest
from docx import Document
from pydantic import ValidationError

from common.ir import Block, BlockType, IRDocument, SourceFormat, Table, TableCell
from common.supervision import SupervisionExtractRequest
from pipeline.index.object_store import ObjectStore
from pipeline.parsing.light_parser import LightParser
from pipeline.supervision.extract import prepare_extraction
from pipeline.supervision.service import (
    SupervisionConfig,
    SupervisionNotReady,
    SupervisionService,
)


def request_data():
    return {
        "metadata": {
            "documentId": "DOC-1",
            "documentVersionId": "DV-1",
            "fileName": "监管函.docx",
            "title": "上传标题",
            "issueDate": "2026-06-30",
            "categoryCode": "REG",
            "documentOrigin": "external",
            "organizationIds": ["ORG-A", "ORG-B"],
            "uploadEntry": "supervision",
        },
        "organizations": [
            {"organizationId": "ORG-A", "name": "甲单位"},
            {"organizationId": "ORG-B", "name": "乙单位"},
        ],
        "rules": [
            {
                "ruleId": "REG-1",
                "reportSection": "external.regulatory",
                "extractFields": [
                    {"key": "problem", "description": "问题事实", "required": True},
                    {"key": "rectification", "description": "整改情况", "required": True},
                ],
            }
        ],
    }


def make_ir():
    return IRDocument(
        doc_version_id="DV-1",
        source_format=SourceFormat.PDF,
        title="正文标题",
        blocks=[
            Block(index=0, type=BlockType.PARAGRAPH, text="甲单位存在记录不完整问题。", page=1),
            Block(
                index=1,
                type=BlockType.TABLE,
                page=2,
                table=Table(
                    n_rows=3,
                    n_cols=2,
                    cells=[
                        TableCell(row=0, col=0, text="单位"),
                        TableCell(row=0, col=1, text="整改情况"),
                        TableCell(row=1, col=0, text="甲单位"),
                        TableCell(row=1, col=1, text="已补记录"),
                        TableCell(row=2, col=0, text="乙单位"),
                        TableCell(row=2, col=1, text="正在整改"),
                    ],
                ),
            ),
        ],
    )


class Extractor:
    def __init__(self, mutate=None):
        self.inputs = []
        self.mutate = mutate

    def chat_json(self, system, user):
        body = json.loads(user)
        self.inputs.append(body)
        e = body["evidence"][0]
        fact = {
            "ruleId": "REG-1",
            "organizationIds": ["ORG-A"],
            "values": {
                "problem": {
                    "value": "记录不完整",
                    "evidenceId": e["evidenceId"],
                    "quote": e["text"],
                },
            },
        }
        if self.mutate:
            self.mutate(fact)
        return {"facts": [fact]}


def test_preserves_all_paragraphs_and_each_table_row_without_overwriting_metadata():
    req = SupervisionExtractRequest.model_validate(request_data())
    result = prepare_extraction(req, make_ir(), client=None, max_evidence_chars=1000)
    assert result["extractionStatus"] == "EVIDENCE_READY"
    assert result["facts"] == []
    assert len(result["evidence"]) == 3
    assert result["evidence"][1]["locatorValue"] == "block:1:row:1"
    assert "已补记录" in result["evidence"][1]["text"]
    assert "正在整改" not in result["evidence"][1]["text"]
    assert "正在整改" in result["evidence"][2]["text"]
    assert req.metadata.title == "上传标题"
    assert req.metadata.issueDate == "2026-06-30"
    assert result["metadataChecks"][0]["code"] == "TITLE_DIFFERENCE"


def test_gateway_visits_every_evidence_group_and_keeps_candidates_unconfirmed():
    req = SupervisionExtractRequest.model_validate(request_data())
    client = Extractor()
    result = prepare_extraction(req, make_ir(), client=client, max_evidence_chars=70)
    consumed = [e["evidenceId"] for call in client.inputs for e in call["evidence"]]
    assert consumed == [e["evidenceId"] for e in result["evidence"]]
    assert len(client.inputs) > 1
    assert result["extractionStatus"] == "EXTRACTED"
    assert all(f["confirmationStatus"] == "PENDING_REVIEW" for f in result["facts"])
    assert result["facts"][0]["missingRequiredFields"] == ["rectification"]


@pytest.mark.parametrize(
    "mutation",
    [
        lambda f: f.update(ruleId="UNKNOWN"),
        lambda f: f.update(organizationIds=["ORG-OUTSIDE"]),
        lambda f: f["values"]["problem"].update(evidenceId="FOREIGN-VERSION"),
        lambda f: f["values"]["problem"].update(quote="原文不存在的摘录"),
        lambda f: f["values"].update(unknown=f["values"]["problem"]),
    ],
)
def test_rejects_invented_or_cross_scope_model_output(mutation):
    with pytest.raises(ValueError):
        prepare_extraction(
            SupervisionExtractRequest.model_validate(request_data()),
            make_ir(),
            client=Extractor(mutation),
            max_evidence_chars=1000,
        )


def test_rejects_ir_version_mismatch_and_oversize_instead_of_truncating():
    req = SupervisionExtractRequest.model_validate(request_data())
    ir = make_ir()
    ir.doc_version_id = "WRONG"
    with pytest.raises(ValueError, match="version"):
        prepare_extraction(req, ir, client=None, max_evidence_chars=1000)
    with pytest.raises(ValueError, match="capacity"):
        prepare_extraction(req, make_ir(), client=Extractor(), max_evidence_chars=5)


def test_request_rejects_identity_injection_and_outside_directory():
    raw = request_data()
    raw["metadata"]["sourceRecordId"] = "not-needed"
    with pytest.raises(ValidationError):
        SupervisionExtractRequest.model_validate(raw)
    raw = request_data()
    raw["organizations"][0]["organizationId"] = "WRONG"
    with pytest.raises(ValidationError):
        SupervisionExtractRequest.model_validate(raw)


class Pg:
    def __init__(self):
        self.version = SimpleNamespace(
            logical_id="DOC-1",
            pipeline_status="INDEXED",
            version_status="effective",
            degraded=False,
            ir_object_key="ir/DV-1.json",
        )
        self.chunks = [
            SimpleNamespace(
                chunk_id="CHUNK-1",
                text="正文",
                page_start=1,
                page_end=2,
                chunk_status="effective",
                degraded=False,
            )
        ]

    @contextmanager
    def session(self):
        yield self

    def get(self, model, pk):
        return self.version

    def scalars(self, statement):
        # 断言生产查询限定本次文档版本；不使用真实数据库。
        assert statement.compile().params["doc_version_id_1"] == "DV-1"
        return self.chunks


def test_service_reuses_real_docx_parser_and_objectstore(tmp_path):
    doc = Document()
    doc.add_paragraph("正文标题")
    doc.add_paragraph("甲单位存在记录不完整问题。")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table.rows[0].cells, ["单位", "整改情况"], strict=True):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["甲单位", "已补记录"], strict=True):
        cell.text = text
    buf = io.BytesIO()
    doc.save(buf)
    parsed = LightParser().parse(buf.getvalue(), "docx", scanned_char_per_page_max=50)
    assert parsed.ok
    ir = IRDocument(
        doc_version_id="DV-1",
        source_format=SourceFormat.DOCX,
        title=parsed.title,
        blocks=parsed.blocks,
    )
    store = ObjectStore(tmp_path)
    store.put_ir(ir)
    before = store.get("ir/DV-1.json")
    svc = SupervisionService(
        Pg(), store, SupervisionConfig(backend="none", max_evidence_chars=1000)
    )
    result = svc.extract(SupervisionExtractRequest.model_validate(request_data()))
    assert store.get("ir/DV-1.json") == before
    assert result["uploadedMaterial"]["processingStatus"] == "indexed"
    assert result["uploadedMaterial"]["title"] == "上传标题"
    assert result["uploadedMaterial"]["parseVersion"].startswith("ir-sha256:")
    assert any("已补记录" in e["text"] for e in result["evidence"])


@pytest.mark.parametrize(
    "field,value",
    [
        ("logical_id", "WRONG"),
        ("pipeline_status", "META_REVIEW"),
        ("version_status", "superseded"),
        ("degraded", True),
        ("ir_object_key", None),
    ],
)
def test_service_rejects_unready_versions_before_reading_files(tmp_path, field, value):
    pg = Pg()
    setattr(pg.version, field, value)
    svc = SupervisionService(
        pg,
        ObjectStore(tmp_path),
        SupervisionConfig(
            backend="none",
            max_evidence_chars=1000,
        ),
    )
    with pytest.raises(SupervisionNotReady):
        svc.extract(SupervisionExtractRequest.model_validate(request_data()))


def test_service_rejects_staging_chunks(tmp_path):
    pg = Pg()
    pg.chunks[0].chunk_status = "staging"
    svc = SupervisionService(
        pg,
        ObjectStore(tmp_path),
        SupervisionConfig(
            backend="none",
            max_evidence_chars=1000,
        ),
    )
    with pytest.raises(SupervisionNotReady, match="index"):
        svc.extract(SupervisionExtractRequest.model_validate(request_data()))


def test_model_call_cannot_return_a_snapshot_changed_during_extraction(tmp_path):
    store = ObjectStore(tmp_path)
    store.put_ir(make_ir())
    pg = Pg()
    client = Extractor(lambda _fact: setattr(pg.chunks[0], "text", "内容已变化"))
    svc = SupervisionService(
        pg,
        store,
        SupervisionConfig(
            backend="gateway",
            max_evidence_chars=1000,
        ),
        client,
    )
    with pytest.raises(SupervisionNotReady, match="changed"):
        svc.extract(SupervisionExtractRequest.model_validate(request_data()))


def test_invalid_business_date_is_reported_not_replaced():
    raw = request_data()
    raw["metadata"]["issueDate"] = "2026-02-30"
    req = SupervisionExtractRequest.model_validate(raw)
    result = prepare_extraction(req, make_ir(), client=None, max_evidence_chars=1000)
    assert {"field": "issueDate", "code": "INVALID_ISSUE_DATE"} in result["metadataChecks"]
    assert req.metadata.issueDate == "2026-02-30"


def test_pdf_line_windows_preserve_exact_text_pages_and_table_boundaries():
    req = SupervisionExtractRequest.model_validate(request_data())
    ir = make_ir()
    ir.blocks = [
        Block(index=0, type=BlockType.PARAGRAPH, text="甲单位应在收到决定书后", page=1),
        Block(index=1, type=BlockType.PARAGRAPH, text="20 个交易日内提交整改报告。", page=1),
        ir.blocks[1].model_copy(update={"index": 2, "page": 1}),
        Block(index=3, type=BlockType.PARAGRAPH, text="下一段", page=1),
        Block(index=4, type=BlockType.PARAGRAPH, text="下一页", page=2),
    ]
    before = ir.model_dump_json()
    result = prepare_extraction(req, ir, client=None, max_evidence_chars=1000)
    window = result["evidence"][0]
    assert window["locatorType"] == "PAGE_TEXT"
    assert window["locatorValue"] == "blocks:0-1"
    assert window["text"] == "甲单位应在收到决定书后\n20 个交易日内提交整改报告。"
    assert [e["locatorType"] for e in result["evidence"]] == [
        "PAGE_TEXT", "TABLE_ROW", "TABLE_ROW", "PARAGRAPH", "PARAGRAPH"
    ]
    assert result["evidence"][-1]["pageStart"] == 2
    assert ir.model_dump_json() == before


def test_location_is_derived_and_cross_line_quotes_require_exact_whitespace():
    raw = request_data()
    raw["rules"][0]["extractFields"].append(
        {"key": "evidenceLocation", "description": "定位", "required": True}
    )
    req = SupervisionExtractRequest.model_validate(raw)
    ir = make_ir()
    ir.blocks = [
        Block(index=0, type=BlockType.PARAGRAPH, text="甲单位存在", page=3),
        Block(index=1, type=BlockType.PARAGRAPH, text="记录不完整问题。", page=3),
    ]
    client = Extractor()
    result = prepare_extraction(req, ir, client=client, max_evidence_chars=1000)
    sent_fields = client.inputs[0]["rules"][0]["extractFields"]
    assert all(f["key"] != "evidenceLocation" for f in sent_fields)
    fact = result["facts"][0]
    assert fact["values"]["evidenceLocation"]["value"] == "第3页；blocks:0-1"
    assert "evidenceLocation" not in fact["missingRequiredFields"]
    restored = prepare_extraction(req, ir, max_evidence_chars=1000, client=Extractor(
        lambda f: f["values"]["problem"].update(quote="甲单位存在记录不完整问题。")
    ))
    assert restored["facts"][0]["values"]["problem"]["quote"] == "甲单位存在\n记录不完整问题。"
    assert any(c.get("code") == "PDF_QUOTE_LINE_BREAKS_RESTORED"
               for c in restored["metadataChecks"])
    with pytest.raises(ValueError, match="non-verbatim"):
        prepare_extraction(req, ir, max_evidence_chars=1000, client=Extractor(
            lambda f: f["values"]["problem"].update(quote="甲单位存在 记录不完整问题。")
        ))
    with pytest.raises(ValueError, match="unknown evidence"):
        prepare_extraction(req, ir, max_evidence_chars=1000, client=Extractor(
            lambda f: f["values"]["problem"].update(evidenceId="DV-1:blocks:0-1")
        ))


class CrossPageExtractor:
    def __init__(self, mutate=None):
        self.mutate = mutate

    def chat_json(self, system, user):
        evidence = json.loads(user)["evidence"]
        value = {
            "value": "甲单位在申报后才结清代持款且未披露；甲单位保荐核查不到位。",
            "evidenceId": evidence[0]["evidenceId"],
            "quote": evidence[0]["text"],
            "supportingEvidence": [{
                "evidenceId": evidence[1]["evidenceId"], "quote": evidence[1]["text"]
            }],
        }
        if self.mutate:
            self.mutate(value)
        return {"facts": [{"ruleId": "REG-1", "organizationIds": ["ORG-A"],
                           "values": {"problem": value}}]}


def cross_page_ir():
    return IRDocument(
        doc_version_id="DV-1", source_format=SourceFormat.PDF,
        blocks=[
            Block(index=0, type=BlockType.PARAGRAPH, page=2,
                  text="甲单位在申报后才结清代持款且未披露。"),
            Block(index=1, type=BlockType.PARAGRAPH, page=3,
                  text="甲单位保荐核查不到位。"),
        ],
    )


def test_cross_page_fact_keeps_details_both_citations_and_derived_locations():
    raw = request_data()
    raw["rules"][0]["extractFields"].append(
        {"key": "evidenceLocation", "description": "定位", "required": True}
    )
    result = prepare_extraction(SupervisionExtractRequest.model_validate(raw), cross_page_ir(),
                                client=CrossPageExtractor(), max_evidence_chars=1000)
    assert len(result["facts"]) == 1
    fact = result["facts"][0]
    value = fact["values"]["problem"]
    assert "未披露" in value["value"] and "核查不到位" in value["value"]
    assert value["quote"] == "甲单位在申报后才结清代持款且未披露。"
    assert value["supportingEvidence"][0]["quote"] == "甲单位保荐核查不到位。"
    location = fact["values"]["evidenceLocation"]
    assert location["value"] == "第2页；block:0 / 第3页；block:1"
    assert location["supportingEvidence"] == value["supportingEvidence"]
    assert fact["confirmationStatus"] == "PENDING_REVIEW"


@pytest.mark.parametrize("mutation", [
    lambda v: v["supportingEvidence"][0].update(evidenceId="ANOTHER-VERSION:block:1"),
    lambda v: v["supportingEvidence"][0].update(quote="甲单位已完成整改。"),
    lambda v: v["supportingEvidence"][0].update(quote="甲单位 保荐核查不到位。"),
])
def test_rejects_invalid_supporting_evidence(mutation):
    with pytest.raises(ValueError, match="unknown evidence or non-verbatim"):
        prepare_extraction(
            SupervisionExtractRequest.model_validate(request_data()), cross_page_ir(),
            client=CrossPageExtractor(mutation), max_evidence_chars=1000,
        )


def test_supporting_evidence_cannot_reference_another_model_group():
    req = SupervisionExtractRequest.model_validate(request_data())
    ir = cross_page_ir()
    evidence = prepare_extraction(req, ir, client=None, max_evidence_chars=1000)["evidence"]
    client = Extractor(lambda f: f["values"]["problem"].update(supportingEvidence=[{
        "evidenceId": evidence[1]["evidenceId"], "quote": evidence[1]["text"]
    }]))
    with pytest.raises(ValueError, match="unknown evidence"):
        prepare_extraction(req, ir, client=client, max_evidence_chars=len(evidence[0]["text"]))


def test_legacy_single_citation_keeps_shape_and_fact_hash():
    req = SupervisionExtractRequest.model_validate(request_data())
    result = prepare_extraction(req, make_ir(), client=Extractor(), max_evidence_chars=1000)
    evidence = result["evidence"][0]
    legacy = {"ruleId": "REG-1", "organizationIds": ["ORG-A"], "values": {"problem": {
        "value": "记录不完整", "evidenceId": evidence["evidenceId"], "quote": evidence["text"]
    }}}
    fact = result["facts"][0]
    assert fact["values"] == legacy["values"]
    assert fact["factId"] == hashlib.sha256(
        json.dumps(legacy, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()


@pytest.mark.parametrize("lines,quote", [
    (["甲单位", "未披露。"], "甲单位已披露。"),
    (["甲单位", "未披露。"], "甲单位 未披露。"),
    (["甲单位", "未披露。"], "甲单位未\n披露" + "新增事项。"),
    (["甲单位", "未披露。甲单位", "未披露。"], "甲单位未披露。"),
    (["收到决定书后 20", "个交易日。"], "收到决定书后 2\n个交易日。"),
])
def test_pdf_quote_alignment_rejects_changes_and_ambiguous_matches(lines, quote):
    ir = make_ir()
    ir.blocks = [Block(index=i, type=BlockType.PARAGRAPH, text=line, page=1)
                 for i, line in enumerate(lines)]
    with pytest.raises(ValueError, match="non-verbatim"):
        prepare_extraction(SupervisionExtractRequest.model_validate(request_data()), ir,
                           client=Extractor(lambda f: f["values"]["problem"].update(quote=quote)),
                           max_evidence_chars=1000)


def test_pdf_supporting_quote_is_restored_without_changing_ir():
    ir = cross_page_ir()
    ir.blocks = [ir.blocks[0],
                 Block(index=1, type=BlockType.PARAGRAPH, text="甲单位保荐", page=3),
                 Block(index=2, type=BlockType.PARAGRAPH, text="核查不到位。", page=3)]
    before = ir.model_dump_json()
    client = CrossPageExtractor(lambda v: v["supportingEvidence"][0].update(
        quote="甲单位保荐核查\n不到位。"
    ))
    result = prepare_extraction(SupervisionExtractRequest.model_validate(request_data()), ir,
                                client=client, max_evidence_chars=1000)
    quote = result["facts"][0]["values"]["problem"]["supportingEvidence"][0]["quote"]
    assert quote == "甲单位保荐\n核查不到位。"
    assert ir.model_dump_json() == before


def test_docx_quote_line_breaks_are_not_rewritten():
    ir = cross_page_ir()
    ir.source_format = SourceFormat.DOCX
    ir.blocks[1].text = "甲单位保荐\n核查不到位。"
    client = CrossPageExtractor(lambda v: v["supportingEvidence"][0].update(
        quote="甲单位保荐核查不到位。"
    ))
    with pytest.raises(ValueError, match="non-verbatim"):
        prepare_extraction(SupervisionExtractRequest.model_validate(request_data()), ir,
                           client=client, max_evidence_chars=1000)
